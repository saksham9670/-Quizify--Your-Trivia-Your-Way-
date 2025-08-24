import os
from flask import Flask, render_template, request, send_file, session, jsonify
import pdfplumber
from docx import Document
from werkzeug.utils import secure_filename
import google.generativeai as genai
from fpdf import FPDF

# Set your API key
os.environ["GOOGLE_API_KEY"] = "AIzaSyCVO9-cJydNCqcQZj1lF2HiSiw3mjV6CDw"
genai.configure(api_key=os.environ["GOOGLE_API_KEY"])
model = genai.GenerativeModel("models/gemini-1.5-flash")

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads/'
app.config['RESULTS_FOLDER'] = 'results/'
app.config['ALLOWED_EXTENSIONS'] = {'pdf', 'txt', 'docx'}
app.secret_key = 'your_secret_key'

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in app.config['ALLOWED_EXTENSIONS']

def extract_text_from_file(file_path):
    ext = file_path.rsplit('.', 1)[1].lower()
    if ext == 'pdf':
        with pdfplumber.open(file_path) as pdf:
            text = ''.join([page.extract_text() for page in pdf.pages])
        return text
    elif ext == 'docx':
        doc = Document(file_path)
        text = ' '.join([para.text for para in doc.paragraphs])
        return text
    elif ext == 'txt':
        with open(file_path, 'r') as file:
            return file.read()
    return None

def Question_mcqs_generator(input_text, num_questions):
    prompt = f"""
    You are an AI assistant helping the user generate multiple-choice questions (MCQs) based on the following text or topic:
    '{input_text}'
    Please generate {num_questions} MCQs. Each question should have:
    - A clear question
    - Four answer options (labeled A, B, C, D)
    - The correct answer clearly indicated
    Format:
    ## MCQ
    Question: [question]
    A) [option A]
    B) [option B]
    C) [option C]
    D) [option D]
    Correct Answer: [correct option]
    """
    response = model.generate_content(prompt).text.strip()
    return response

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/generate', methods=['POST'])
def generate_mcqs():
    method = request.form.get('method')

    if method == 'document':
        file = request.files.get('file')
        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            file.save(file_path)
            text = extract_text_from_file(file_path)
            if text:
                num_questions = int(request.form['num_questions'])
                mcqs = Question_mcqs_generator(text, num_questions)
                session['mcqs'] = [q.strip() for q in mcqs.split("## MCQ") if q.strip()]  # Filter out empty questions
                return render_template('game.html', total=len(session['mcqs']))
    elif method == 'topic':
        topic = request.form.get('topic')
        num_questions = int(request.form['num_questions'])
        mcqs = Question_mcqs_generator(topic, num_questions)
        session['mcqs'] = [q.strip() for q in mcqs.split("## MCQ") if q.strip()]  # Filter out empty questions
        return render_template('game.html', total=len(session['mcqs']))

    return "Invalid request"
@app.route('/get-question/<int:index>', methods=['GET'])
def get_question(index):
    mcqs = session.get('mcqs', [])
    if 0 <= index < len(mcqs):
        question = mcqs[index].strip()
        question_text = question.split('A)')[0].strip()
        options = {
            'A': question.split('A)')[1].split('B)')[0].strip(),
            'B': question.split('B)')[1].split('C)')[0].strip(),
            'C': question.split('C)')[1].split('D)')[0].strip(),
            'D': question.split('D)')[1].split('Correct Answer:')[0].strip()
        }
        correct_answer = question.split('Correct Answer:')[1].strip()
        
        return jsonify({
            'question': question_text,
            'options': options,
            'correct': correct_answer
        })
    return jsonify({'error': 'Invalid index'}), 400


@app.route('/results')
def show_results():
    score = session.get('score', 0)
    total = len(session.get('mcqs', []))
    return render_template('results.html', score=score, total=total)

@app.route('/submit-score', methods=['POST'])
def submit_score():
    session['score'] = request.json.get('score', 0)
    return jsonify({'message': 'Score saved successfully'})


@app.route('/download-pdf')
def download_pdf():
    mcqs = session.get('mcqs', [])
    if not mcqs:
        return "No MCQs to download"

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font('Arial', 'B', 16)
    pdf.cell(200, 10, txt="MCQ Questions and Answers", ln=True, align='C')

    for i, mcq in enumerate(mcqs):
        try:
            # Extract question and options
            parts = mcq.split('A)')
            question = parts[0].strip()  # Question text
            options = parts[1].split('B)')
            option_a = options[0].strip()
            options = options[1].split('C)')
            option_b = options[0].strip()
            options = options[1].split('D)')
            option_c = options[0].strip()
            option_d = options[1].split('Correct Answer:')[0].strip()
            correct_answer = options[1].split('Correct Answer:')[1].strip()

            # Add question and options to PDF
            pdf.ln(10)
            pdf.set_font('Arial', '', 12)
            pdf.multi_cell(0, 10, f"{i + 1}. {question}")
            pdf.multi_cell(0, 10, f"  A) {option_a}")
            pdf.multi_cell(0, 10, f"  B) {option_b}")
            pdf.multi_cell(0, 10, f"  C) {option_c}")
            pdf.multi_cell(0, 10, f"  D) {option_d}")
            pdf.multi_cell(0, 10, f"Correct Answer: {correct_answer}")
        except Exception as e:
            pdf.ln(10)
            pdf.set_font('Arial', 'I', 12)
            pdf.multi_cell(0, 10, f"Error processing MCQ {i + 1}: {str(e)}")

    # Save and serve PDF
    if not os.path.exists(app.config['RESULTS_FOLDER']):
        os.makedirs(app.config['RESULTS_FOLDER'])
    pdf_output_path = os.path.join(app.config['RESULTS_FOLDER'], 'mcq_questions.pdf')
    pdf.output(pdf_output_path)

    return send_file(pdf_output_path, as_attachment=True)


if __name__ == "__main__":
    app.run(debug=True)
